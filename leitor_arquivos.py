"""Leitura de arquivo de DRE/Balancete enviado pelo usuário — PDF ou
Excel. Extrai texto/células e tenta achar, por nome de conta (busca de
texto, não posição fixa), os valores que o formulário do dashboard
pede.

**Por que isso é uma SUGESTÃO, não preenchimento automático confiável
de verdade:** a busca é por rótulo de conta num texto solto (ainda
mais no caso de PDF — a ordem do texto extraído nem sempre segue a
ordem visual das colunas), nunca vi a estrutura real de um DRE/
Balancete exportado deste escritório especificamente. Pode achar o
número errado (conta parecida, linha de total em vez de linha
detalhada, etc.). Por isso cada campo encontrado é marcado como
"sugestão" no formulário — sempre precisa de conferência humana antes
de gerar o dashboard pro cliente, nunca é usado silenciosamente.
"""

import re
import unicodedata
from pathlib import Path

PASTA_UPLOADS = Path(__file__).resolve().parent / "data" / "uploads"
EXTENSOES_ACEITAS = {".pdf", ".xlsx", ".xls"}
TAMANHO_MAXIMO_BYTES = 15 * 1024 * 1024  # 15 MB — generoso pra um DRE/Balancete, barra arquivo errado


def salvar_upload(arquivo_flask):
    """Salva o arquivo enviado (`werkzeug.FileStorage`) em
    `data/uploads/`, com nome único (timestamp + nome original) pra
    não sobrescrever um upload anterior. Devolve o `Path` salvo, ou
    `None` se a extensão não for aceita."""
    import datetime

    nome_original = arquivo_flask.filename or ""
    extensao = Path(nome_original).suffix.lower()
    if extensao not in EXTENSOES_ACEITAS:
        return None

    PASTA_UPLOADS.mkdir(parents=True, exist_ok=True)
    carimbo = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_seguro = "".join(c for c in Path(nome_original).stem if c.isalnum() or c in "-_") or "arquivo"
    caminho = PASTA_UPLOADS / f"{carimbo}_{nome_seguro}{extensao}"
    arquivo_flask.save(caminho)
    return caminho


def extrair_conteudo(caminho):
    """Devolve `(texto_extraido, metodo, erro)` — `texto_extraido` é
    `None` só se a extração falhar por completo. Nunca lança exceção
    pro chamador (arquivo corrompido/formato inesperado vira erro
    textual, não um crash da página)."""
    caminho = Path(caminho)
    extensao = caminho.suffix.lower()

    if extensao == ".pdf":
        return _extrair_pdf(caminho)
    if extensao in (".xlsx", ".xls"):
        return _extrair_excel(caminho)
    return None, None, f"Extensão não suportada: {extensao}"


def _extrair_pdf(caminho):
    try:
        import pypdf

        leitor = pypdf.PdfReader(str(caminho))
        partes = [pagina.extract_text() or "" for pagina in leitor.pages]
        texto = "\n".join(partes).strip()
        if not texto:
            return None, None, "PDF aberto, mas sem texto extraível (pode ser um PDF escaneado/imagem)."
        return texto, "pypdf", None
    except Exception as erro:
        return None, None, f"Não consegui abrir o PDF: {erro!r}"


# Campo do formulário -> variações de nome de conta que podem
# aparecer no DRE/Balancete real (ordem importa: a primeira variação
# que bater primeiro no texto é a usada). Lista enxuta de propósito —
# melhor não achar nada do que achar errado por um sinônimo vago
# demais.
_ROTULOS_POR_CAMPO = {
    "receita_bruta": ["receita bruta de vendas", "receita operacional bruta", "receita bruta"],
    "deducoes": ["deducoes da receita bruta", "deducoes de vendas", "impostos sobre vendas", "deducoes da receita", "deducoes"],
    "cmv": ["custo das mercadorias vendidas", "custo dos produtos vendidos", "custo dos servicos prestados", "custo das vendas", "cmv", "cpv"],
    "despesas_operacionais": ["total das despesas operacionais", "despesas operacionais", "despesas administrativas e comerciais"],
    "resultado_financeiro": ["resultado financeiro liquido", "resultado financeiro"],
    "irpj_csll": ["provisao para irpj e csll", "irpj e csll", "irpj/csll", "imposto de renda e contribuicao social"],
    "ativo_circulante": ["total do ativo circulante", "ativo circulante"],
    "ativo_nao_circulante": ["total do ativo nao circulante", "ativo nao circulante", "ativo realizavel a longo prazo"],
    "passivo_circulante": ["total do passivo circulante", "passivo circulante"],
    "passivo_nao_circulante": ["total do passivo nao circulante", "passivo nao circulante", "exigivel a longo prazo"],
    "patrimonio_liquido": ["total do patrimonio liquido", "patrimonio liquido"],
}

_PADRAO_NUMERO = re.compile(r"\(?-?\s*(?:R\$\s*)?\d{1,3}(?:\.\d{3})*,\d{2}\)?|\(?-?\s*\d+,\d{2}\)?")


def _sem_acento_minusculo(texto):
    sem_acento = "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )
    return sem_acento.lower()


def _numero_para_float(texto_numero):
    """Converte '1.234,56' ou '(1.234,56)' (negativo entre parênteses,
    formato contábil comum) pra float."""
    negativo = "(" in texto_numero or texto_numero.strip().startswith("-")
    limpo = texto_numero.replace("R$", "").replace("(", "").replace(")", "").replace("-", "").strip()
    limpo = limpo.replace(".", "").replace(",", ".")
    try:
        valor = float(limpo)
    except ValueError:
        return None
    return -valor if negativo else valor


def sugerir_campos(texto_extraido):
    """Procura, linha por linha do texto extraído, os rótulos de conta
    de `_ROTULOS_POR_CAMPO` — quando acha um, pega o ÚLTIMO número da
    mesma linha (relatórios contábeis normalmente põem o valor no fim
    da linha) como valor sugerido pro campo do formulário.

    Devolve um dict `{campo: {"valor": float, "linha": str} }` só com
    os campos que acharam alguma coisa — campo ausente no dict =
    "não achei, preencha à mão". Valores de conta redutora (deduções,
    CMV, despesas, IRPJ/CSLL) são devolvidos em módulo (positivo) —
    o cálculo em `dashboard.py` já assume isso e subtrai internamente.
    """
    linhas = [l for l in texto_extraido.splitlines() if l.strip()]
    linhas_normalizadas = [_sem_acento_minusculo(l) for l in linhas]

    sugestoes = {}
    for campo, rotulos in _ROTULOS_POR_CAMPO.items():
        for rotulo in rotulos:
            achou = False
            for linha_original, linha_norm in zip(linhas, linhas_normalizadas):
                if rotulo in linha_norm:
                    numeros = _PADRAO_NUMERO.findall(linha_original)
                    if not numeros:
                        continue
                    valor = _numero_para_float(numeros[-1])
                    if valor is None:
                        continue
                    # Todo campo é uma magnitude positiva (dashboard.py subtrai
                    # internamente), EXCETO resultado_financeiro — esse pode
                    # ser negativo de verdade (despesa financeira líquida) e
                    # precisa manter o sinal que veio do arquivo.
                    valor_final = valor if campo == "resultado_financeiro" else abs(valor)
                    sugestoes[campo] = {"valor": valor_final, "linha": linha_original.strip()}
                    achou = True
                    break
            if achou:
                break  # já achou esse campo com a variação de rótulo mais específica; não tenta as outras

    return sugestoes


def _extrair_excel(caminho):
    try:
        import openpyxl

        pasta = openpyxl.load_workbook(caminho, data_only=True)
        partes = []
        for nome_aba in pasta.sheetnames:
            aba = pasta[nome_aba]
            partes.append(f"--- Aba: {nome_aba} ---")
            for linha in aba.iter_rows(values_only=True):
                valores = [str(v) for v in linha if v is not None]
                if valores:
                    partes.append(" | ".join(valores))
        texto = "\n".join(partes).strip()
        if not texto:
            return None, None, "Planilha aberta, mas não achei nenhuma célula com conteúdo."
        return texto, "openpyxl", None
    except Exception as erro:
        return None, None, f"Não consegui abrir a planilha: {erro!r}"
